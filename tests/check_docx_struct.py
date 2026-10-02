import docx

doc = docx.Document('Approved Judgment_Notarised+Legalised_SN 193344 (1)-Non-Parsable-en-US#SRECFMT_DXBRBK#.docx')

print(f"Total sections: {len(doc.sections)}")
for i, s in enumerate(doc.sections):
    h_text = [p.text for p in s.header.paragraphs if p.text.strip()]
    f_text = [p.text for p in s.footer.paragraphs if p.text.strip()]
    print(f"Section {i} Header: {h_text}")
    print(f"Section {i} Footer: {f_text}")

breaks = []
for i, p in enumerate(doc.paragraphs):
    for r in p.runs:
        if 'lastRenderedPageBreak' in r._r.xml or ('w:br' in r._r.xml and 'page' in r._r.xml):
            breaks.append((i+1, p.text[:50]))

print(f"Page breaks count: {len(breaks)}")
for b in breaks:
    print(f"  P{b[0]}: {b[1]}")
