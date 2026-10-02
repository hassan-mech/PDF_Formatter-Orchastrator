"""
check_docx.py — DOCX integrity checker: XML validity, suspicious characters, fragmented runs.

Replaces: check_docx_struct.py + check_chars.py

Usage:
    python check_docx.py -i output/document.docx --check all
    python check_docx.py -i output/document.docx --check chars
    python check_docx.py -i output/document.docx --check runs
"""

import argparse
import sys
import unicodedata
import zipfile
from pathlib import Path
from xml.etree import ElementTree as ET

try:
    from docx import Document
    from docx.oxml.ns import qn
except ImportError:
    print("[ERROR] python-docx not installed. Run: pip install python-docx")
    sys.exit(1)


# ---------------------------------------------------------------------------
# Check: XML well-formedness
# ---------------------------------------------------------------------------

def check_xml(docx_path: Path) -> int:
    """Check that all XML parts inside the DOCX zip are well-formed. Returns error count."""
    print("\n📋 XML Well-Formedness Check")
    print("  " + "-" * 50)
    errors = 0
    with zipfile.ZipFile(str(docx_path), "r") as z:
        xml_files = [n for n in z.namelist() if n.endswith(".xml") or n.endswith(".rels")]
        for name in sorted(xml_files):
            try:
                data = z.read(name)
                ET.fromstring(data)
                print(f"  ✅ {name}")
            except ET.ParseError as exc:
                print(f"  ❌ {name}: {exc}")
                errors += 1
    print(f"\n  Result: {len(xml_files)-errors}/{len(xml_files)} XML parts valid, {errors} error(s)")
    return errors


# ---------------------------------------------------------------------------
# Check: Suspicious / non-standard Unicode characters
# ---------------------------------------------------------------------------

# Categories considered suspicious for a typical Latin/Arabic document
_SUSPICIOUS_CATEGORIES = {
    "Cc",  # Control characters
    "Cs",  # Surrogate characters
    "Co",  # Private use
    "Cn",  # Unassigned
}

_ALLOWED_CONTROL = {"\t", "\n", "\r"}  # Common whitespace controls that are fine


def is_suspicious(ch: str) -> bool:
    """Return True for characters that are likely erroneous or non-standard."""
    if ch in _ALLOWED_CONTROL:
        return False
    cat = unicodedata.category(ch)
    if cat in _SUSPICIOUS_CATEGORIES:
        return True
    # Flag non-breaking / zero-width spaces and directional marks
    code = ord(ch)
    if code in (0x00A0, 0x200B, 0x200C, 0x200D, 0x200E, 0x200F, 0xFEFF, 0xFFFD):
        return True
    return False


def check_chars(docx_path: Path) -> int:
    """Scan all paragraph text for suspicious Unicode characters. Returns total count found."""
    print("\n🔤 Suspicious Character Check")
    print("  " + "-" * 50)
    doc = Document(str(docx_path))
    total_found = 0
    report = []

    all_paragraphs = list(doc.paragraphs)
    for tbl in doc.tables:
        for row in tbl.rows:
            for cell in row.cells:
                all_paragraphs.extend(cell.paragraphs)

    for p_idx, para in enumerate(all_paragraphs):
        text = para.text
        found = []
        for c_idx, ch in enumerate(text):
            if is_suspicious(ch):
                found.append((c_idx, ch, ord(ch), unicodedata.name(ch, "?")))
        if found:
            total_found += len(found)
            report.append((p_idx + 1, text[:60], found))

    if report:
        for p_num, sample, chars in report:
            print(f"  ⚠️  Para {p_num} [{repr(sample)}]:")
            for c_idx, ch, code, name in chars:
                print(f"       pos={c_idx}  U+{code:04X}  [{name}]  repr={repr(ch)}")
    else:
        print("  ✅ No suspicious characters found.")

    print(f"\n  Result: {total_found} suspicious character(s) across {len(report)} paragraph(s)")
    return total_found


# ---------------------------------------------------------------------------
# Check: Fragmented runs
# ---------------------------------------------------------------------------

def _runs_same_format(r1, r2) -> bool:
    """Return True if two docx Run objects have identical rPr XML."""
    from lxml import etree
    rpr1 = r1._r.find(qn("w:rPr"))
    rpr2 = r2._r.find(qn("w:rPr"))
    if rpr1 is None and rpr2 is None:
        return True
    if rpr1 is None or rpr2 is None:
        return False
    return etree.tostring(rpr1) == etree.tostring(rpr2)


def check_runs(docx_path: Path) -> int:
    """Report paragraphs with adjacent runs that share identical formatting. Returns frag count."""
    print("\n🧩 Fragmented Runs Check")
    print("  " + "-" * 50)
    doc = Document(str(docx_path))
    total_frag = 0
    frag_paras = 0

    all_paragraphs = list(doc.paragraphs)
    for tbl in doc.tables:
        for row in tbl.rows:
            for cell in row.cells:
                all_paragraphs.extend(cell.paragraphs)

    for p_idx, para in enumerate(all_paragraphs):
        runs = para.runs
        if len(runs) < 2:
            continue
        frag_count = 0
        for i in range(len(runs) - 1):
            if _runs_same_format(runs[i], runs[i + 1]):
                frag_count += 1
        if frag_count > 0:
            total_frag += frag_count
            frag_paras += 1
            sample = para.text[:60]
            print(
                f"  ⚠️  Para {p_idx+1}: {frag_count} mergeable run pair(s)  "
                f"[{len(runs)} runs]  [{repr(sample)}]"
            )

    if frag_paras == 0:
        print("  ✅ No fragmented runs found.")
    print(f"\n  Result: {total_frag} fragmented run pair(s) in {frag_paras} paragraph(s)")
    return total_frag


# ---------------------------------------------------------------------------
# Main
# ---------------------------------------------------------------------------

def main():
    parser = argparse.ArgumentParser(
        description="Check DOCX integrity: XML, suspicious chars, fragmented runs."
    )
    parser.add_argument("-i", "--input", required=True, help="Path to .docx file")
    parser.add_argument(
        "--check",
        choices=["xml", "chars", "runs", "all"],
        default="all",
        help="Which check(s) to run (default: all)"
    )
    args = parser.parse_args()

    docx_path = Path(args.input)
    if not docx_path.exists():
        print(f"[ERROR] File not found: {docx_path}")
        sys.exit(1)

    print(f"🔍 Checking DOCX: {docx_path.name}")
    issues = 0

    if args.check in ("xml", "all"):
        issues += check_xml(docx_path)

    if args.check in ("chars", "all"):
        issues += check_chars(docx_path)

    if args.check in ("runs", "all"):
        issues += check_runs(docx_path)

    print(f"\n{'='*55}")
    overall = "✅ PASS" if issues == 0 else f"⚠️  {issues} issue(s) found"
    print(f"  Overall: {overall}")
    print(f"{'='*55}")
    sys.exit(0 if issues == 0 else 1)


if __name__ == "__main__":
    main()
