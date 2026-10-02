import fitz
import docx as docx_lib
from pathlib import Path

pdf_doc = fitz.open("input/DOC0000074469-Non-Parsable-en-US#FPREP_DXEGTK#.pdf")
pdf_text = "\n".join(pdf_doc[i].get_text("text").strip() for i in range(pdf_doc.page_count))

docx_path = Path("runs/DOC0000074469_20261002_1045/build/DOC0000074469-Non-Parsable-en-US#FPREP_DXEGTK#.docx")
doc = docx_lib.Document(str(docx_path))

parts = [p.text for p in doc.paragraphs if p.text]
seen_cells = set()

def extract_from_table(tbl):
    for row in tbl.rows:
        for cell in row.cells:
            if cell._tc not in seen_cells:
                seen_cells.add(cell._tc)
                for p in cell.paragraphs:
                    if p.text.strip():
                        parts.append(p.text.strip())
                for nested in cell.tables:
                    extract_from_table(nested)

for t in doc.tables:
    extract_from_table(t)

docx_text = "\n".join(parts)

words_pdf = set(pdf_text.lower().split())
words_docx = set(docx_text.lower().split())

missing = words_pdf - words_docx
recall = len(words_pdf & words_docx) / len(words_pdf)
print(f"Token recall with nested tables: {recall:.1%}")
print(f"Missing count: {len(missing)}")
print("Sample missing words:", sorted(list(missing))[:50])
