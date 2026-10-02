import docx
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls, qn
from docx.enum.section import WD_SECTION_START
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_BREAK

doc = docx.Document()
section = doc.sections[0]
sectPr = section._sectPr

# Set 2 equal columns
cols_xml = f'<w:cols {nsdecls("w")} w:num="2" w:space="360" w:equalWidth="1"/>'
sectPr.append(parse_xml(cols_xml))

p1 = doc.add_paragraph('Left column text')
p2 = doc.add_paragraph('More left column text')
r2 = p2.add_run()
r2.add_break(WD_BREAK.COLUMN)
p3 = doc.add_paragraph('Right column text')

doc.save('runs/DOC0000074469_20261002_1045/test_2col.docx')
print('Successfully created test_2col.docx')
