import docx
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls, qn
from docx.enum.section import WD_SECTION_START
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_BREAK

doc = docx.Document()
section = doc.sections[0]
section.page_width = Inches(8.27)
section.page_height = Inches(10.63)
section.top_margin = Inches(0.5)
section.bottom_margin = Inches(0.5)
section.left_margin = Inches(0.98)
section.right_margin = Inches(0.98)

# 1-column title section
p_title = doc.add_paragraph("Melanoma metastásico, un caso extraordinario...")

# Continuous section break for 2 unequal columns
s2 = doc.add_section(WD_SECTION_START.CONTINUOUS)
sectPr = s2._sectPr
cols_xml = f'''<w:cols {nsdecls("w")} w:num="2" w:equalWidth="0">
  <w:col w:w="5900" w:space="360"/>
  <w:col w:w="2600"/>
</w:cols>'''
sectPr.append(parse_xml(cols_xml))

p_left = doc.add_paragraph("Resumen text in wider left column")
r = p_left.add_run()
r.add_break(WD_BREAK.COLUMN)
p_right = doc.add_paragraph("Affiliations text in narrower right column")

doc.save("runs/DOC0000074469_20261002_1045/test_unequal_cols.docx")
print("Saved test_unequal_cols.docx")
