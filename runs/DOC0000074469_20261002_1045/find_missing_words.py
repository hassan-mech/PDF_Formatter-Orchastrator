import fitz
import docx as docx_lib
from pathlib import Path

pdf_doc = fitz.open("input/DOC0000074469-Non-Parsable-en-US#FPREP_DXEGTK#.pdf")
pdf_text = "\n".join(pdf_doc[i].get_text("text").strip() for i in range(pdf_doc.page_count))

docx_path = Path("runs/DOC0000074469_20261002_1045/build/DOC0000074469-Non-Parsable-en-US#FPREP_DXEGTK#.docx")
doc = docx_lib.Document(str(docx_path))
docx_parts = [p.text for p in doc.paragraphs if p.text]
seen_cells = set()
for t in doc.tables:
    for row in t.rows:
        for cell in row.cells:
            if cell._tc not in seen_cells:
                seen_cells.add(cell._tc)
                if cell.text.strip():
                    docx_parts.append(cell.text.strip())

docx_text = "\n".join(docx_parts)

words_pdf = set(pdf_text.lower().split())
words_docx = set(docx_text.lower().split())

missing = words_pdf - words_docx
print(f"Total unique PDF words: {len(words_pdf)}")
print(f"Total unique DOCX words: {len(words_docx)}")
print(f"Missing words count: {len(missing)}")
print("Sample missing words:", sorted(list(missing))[:50])

# Check per page
for i in range(pdf_doc.page_count):
    page_words = set(pdf_doc[i].get_text("text").lower().split())
    page_missing = page_words - words_docx
    print(f"Page {i+1}: {len(page_words)} words, {len(page_missing)} missing ({len(page_missing)/max(len(page_words),1):.1%})")
    if page_missing:
        print(f"   Missing sample: {sorted(list(page_missing))[:10]}")
