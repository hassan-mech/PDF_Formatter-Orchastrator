import docx
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls
from docx.shared import Inches, Pt, RGBColor

doc = docx.Document()
p1 = doc.add_paragraph("First paragraph in pink box")
pPr1 = p1._p.get_or_add_pPr()
pPr1.append(parse_xml(f'<w:shd {nsdecls("w")} w:fill="FCECEF"/>'))
p1.paragraph_format.space_before = Pt(0)
p1.paragraph_format.space_after = Pt(2)

p2 = doc.add_paragraph("Second paragraph in pink box")
pPr2 = p2._p.get_or_add_pPr()
pPr2.append(parse_xml(f'<w:shd {nsdecls("w")} w:fill="FCECEF"/>'))
p2.paragraph_format.space_before = Pt(0)
p2.paragraph_format.space_after = Pt(2)

doc.save("runs/DOC0000074469_20261002_1045/test_shd.docx")
print("Saved test_shd.docx")
