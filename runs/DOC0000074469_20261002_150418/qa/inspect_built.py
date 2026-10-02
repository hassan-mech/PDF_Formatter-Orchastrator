import docx

doc = docx.Document("runs/DOC0000074469_20261002_150418/build/DOC0000074469-Non-Parsable-en-US#FPREP_DXEGTK#.docx")
print("=== SECTION ANALYSIS ===")
print("Total sections:", len(doc.sections))
for i, s in enumerate(doc.sections):
    sectPr = s._sectPr
    cols = sectPr.xpath("./w:cols")
    col_info = [c.attrib for c in cols] if cols else "None"
    type_info = sectPr.xpath("./w:type")
    type_val = [t.attrib for t in type_info] if type_info else "NEW_PAGE (default)"
    footer_p = [p.text for p in s.footer.paragraphs if p.text]
    print(f"Section {i+1}: type={type_val}, cols={col_info}, margins=(L={s.left_margin.inches:.2f}\", R={s.right_margin.inches:.2f}\", T={s.top_margin.inches:.2f}\", B={s.bottom_margin.inches:.2f}\"), footer={footer_p}")

print("\n=== TABLE ANALYSIS ===")
print("Total tables in DOCX:", len(doc.tables))
for i, t in enumerate(doc.tables):
    print(f"Table {i+1}: rows={len(t.rows)}, cols={len(t.columns)}")

print("\n=== HIDDEN RUNS ANALYSIS ===")
hidden_count = 0
for p_idx, p in enumerate(doc.paragraphs):
    hidden_runs = [r for r in p.runs if r.font.hidden]
    if hidden_runs:
        hidden_count += len(hidden_runs)
        print(f"Para {p_idx}: {len(hidden_runs)} hidden run(s), size={hidden_runs[0].font.size.pt}pt, color={hidden_runs[0].font.color.rgb}, text_preview={hidden_runs[0].text[:80]}...")
print("Total hidden runs:", hidden_count)
