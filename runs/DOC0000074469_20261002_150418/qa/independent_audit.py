import docx
from docx.oxml.ns import qn
import pymupdf
from pathlib import Path

docx_path = Path("output/DOC0000074469-Non-Parsable-en-US#FPREP_DXEGTK#.docx")
pdf_path = Path("input/DOC0000074469-Non-Parsable-en-US#FPREP_DXEGTK#.pdf")
if not pdf_path.exists():
    pdf_path = Path("finished/DOC0000074469-Non-Parsable-en-US#FPREP_DXEGTK#.pdf")

doc = docx.Document(str(docx_path))

print("=== ALL 15 SECTIONS DETAILED AUDIT ===")
for i, s in enumerate(doc.sections):
    sectPr = s._sectPr
    type_elem = sectPr.xpath("./w:type")
    start_type = type_elem[0].get(qn("w:val")) if type_elem else "NEW_PAGE"
    
    cols_elem = sectPr.xpath("./w:cols")
    cols_attr = dict(cols_elem[0].attrib) if cols_elem else {}
    col_children = []
    if cols_elem:
        for child in cols_elem[0]:
            if child.tag.endswith("col"):
                col_children.append(dict(child.attrib))
    
    w = s.page_width.inches
    h = s.page_height.inches
    l = round(s.left_margin.inches, 2)
    r = round(s.right_margin.inches, 2)
    t = round(s.top_margin.inches, 2)
    b = round(s.bottom_margin.inches, 2)
    printable_w = round(w - l - r, 2)
    
    footer = s.footer
    f_linked = footer.is_linked_to_previous
    f_texts = [p.text for p in footer.paragraphs if p.text.strip()]
    f_tabs = []
    if footer.paragraphs:
        for ts in footer.paragraphs[0].paragraph_format.tab_stops:
            f_tabs.append(round(ts.position.inches, 2))
            
    header = s.header
    h_linked = header.is_linked_to_previous
    h_texts = [p.text for p in header.paragraphs if p.text.strip()]
    
    cols_desc = f"num={cols_attr.get('{http://schemas.openxmlformats.org/wordprocessingml/2006/main}num', '1')}"
    if cols_attr.get('{http://schemas.openxmlformats.org/wordprocessingml/2006/main}equalWidth') == '0':
        cols_desc += f" (unequal: {col_children})"
    elif cols_attr.get('{http://schemas.openxmlformats.org/wordprocessingml/2006/main}num') == '2':
        cols_desc += f" (equal: space={cols_attr.get('{http://schemas.openxmlformats.org/wordprocessingml/2006/main}space')})"

    print(f"Sec {i+1:2d} | Start: {start_type:10s} | Margins: L={l:.2f}\" R={r:.2f}\" T={t:.2f}\" B={b:.2f}\" | Grid: {printable_w:.2f}\" | Cols: {cols_desc} | Footer (linked={f_linked}, tabs={f_tabs}): {f_texts}")
