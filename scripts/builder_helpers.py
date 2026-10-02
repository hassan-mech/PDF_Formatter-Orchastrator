"""
builder_helpers.py — High-level paragraph/heading builder functions for Word documents.

Provides reusable, parameterised helpers that wrap python-docx to build
consistently formatted paragraphs, headings, numbered lists, and quotes.

Usage:
    from builder_helpers import add_heading, add_numbered_para, add_quote_para, add_simple_para
"""

import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn


# ---------------------------------------------------------------------------
# Internal helpers
# ---------------------------------------------------------------------------

_ALIGN_MAP = {
    "left":    WD_ALIGN_PARAGRAPH.LEFT,
    "right":   WD_ALIGN_PARAGRAPH.RIGHT,
    "center":  WD_ALIGN_PARAGRAPH.CENTER,
    "justify": WD_ALIGN_PARAGRAPH.JUSTIFY,
}


def _apply_keep_with_next(p, keep: bool):
    """Set the keepWithNext paragraph property via OOXML."""
    pPr = p._p.get_or_add_pPr()
    kwn = OxmlElement("w:keepNext")
    if keep:
        pPr.append(kwn)
    else:
        # Remove if present
        for existing in pPr.findall(qn("w:keepNext")):
            pPr.remove(existing)


# ---------------------------------------------------------------------------
# Public builder functions
# ---------------------------------------------------------------------------

def add_heading(
    doc,
    text: str,
    space_before: float = 12,
    space_after: float = 6,
    font_size: float = 11,
    font_name: str = "Times New Roman",
    line_spacing: float = 1.15,
    keep_with_next: bool = True,
) -> "docx.text.paragraph.Paragraph":
    """
    Add an underlined heading paragraph to the document.

    Parameters
    ----------
    doc            : python-docx Document object
    text           : heading text
    space_before   : space above paragraph in points (default 12)
    space_after    : space below paragraph in points (default 6)
    font_size      : font size in points (default 11)
    font_name      : font name (default 'Times New Roman')
    line_spacing   : line spacing multiplier (default 1.15)
    keep_with_next : keep paragraph with the next one (default True)
    """
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(space_before)
    p.paragraph_format.space_after  = Pt(space_after)
    p.paragraph_format.line_spacing = line_spacing
    _apply_keep_with_next(p, keep_with_next)
    r = p.add_run(text)
    r.underline   = True
    r.bold        = False
    r.font.name   = font_name
    r.font.size   = Pt(font_size)
    return p


def add_numbered_para(
    doc,
    num_str: str,
    text_runs: list,
    space_after: float = 6,
    font_size: float = 11,
    font_name: str = "Times New Roman",
    line_spacing: float = 1.15,
    keep_with_next: bool = False,
) -> "docx.text.paragraph.Paragraph":
    """
    Add a numbered paragraph (e.g. '1.' or '(a)') with hanging indent.

    Parameters
    ----------
    doc            : python-docx Document object
    num_str        : number/label string (e.g. '1.', '(a)')
    text_runs      : list of (text, italic, bold, underline) tuples
    space_after    : space below in points (default 6)
    font_size      : font size in points (default 11)
    font_name      : font name (default 'Times New Roman')
    line_spacing   : line spacing multiplier (default 1.15)
    keep_with_next : keep paragraph with next (default False)
    """
    p = doc.add_paragraph()
    p.paragraph_format.left_indent       = Inches(0.4)
    p.paragraph_format.first_line_indent = Inches(-0.4)
    p.paragraph_format.space_after       = Pt(space_after)
    p.paragraph_format.line_spacing      = line_spacing
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    _apply_keep_with_next(p, keep_with_next)

    r_num = p.add_run(f"{num_str}\t")
    r_num.font.name = font_name
    r_num.font.size = Pt(font_size)

    for r_text, r_italic, r_bold, r_underline in text_runs:
        r = p.add_run(r_text)
        r.font.name  = font_name
        r.font.size  = Pt(font_size)
        r.italic     = r_italic
        r.bold       = r_bold
        r.underline  = r_underline
    return p


def add_quote_para(
    doc,
    text_runs: list,
    space_after: float = 6,
    left_indent: float = 0.5,
    font_size: float = 11,
    font_name: str = "Times New Roman",
    line_spacing: float = 1.15,
    keep_with_next: bool = False,
) -> "docx.text.paragraph.Paragraph":
    """
    Add an indented quote paragraph.

    Parameters
    ----------
    doc            : python-docx Document object
    text_runs      : list of (text, italic, bold, underline) tuples
    space_after    : space below in points (default 6)
    left_indent    : left indent in inches (default 0.5)
    font_size      : font size in points (default 11)
    font_name      : font name (default 'Times New Roman')
    line_spacing   : line spacing multiplier (default 1.15)
    keep_with_next : keep paragraph with next (default False)
    """
    p = doc.add_paragraph()
    p.paragraph_format.left_indent  = Inches(left_indent)
    p.paragraph_format.space_after  = Pt(space_after)
    p.paragraph_format.line_spacing = line_spacing
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    _apply_keep_with_next(p, keep_with_next)

    for r_text, r_italic, r_bold, r_underline in text_runs:
        r = p.add_run(r_text)
        r.font.name  = font_name
        r.font.size  = Pt(font_size)
        r.italic     = r_italic
        r.bold       = r_bold
        r.underline  = r_underline
    return p


def add_simple_para(
    doc,
    text: str,
    font_name: str  = "Times New Roman",
    font_size: float = 11,
    bold: bool      = False,
    italic: bool    = False,
    underline: bool = False,
    align: str      = "justify",
    space_before: float = 0,
    space_after: float  = 6,
    left_indent: float  = 0,
    line_spacing: float = 1.15,
) -> "docx.text.paragraph.Paragraph":
    """
    Add a simple single-run paragraph with full formatting control.

    Parameters
    ----------
    doc          : python-docx Document object
    text         : paragraph text
    font_name    : font name (default 'Times New Roman')
    font_size    : font size in points (default 11)
    bold         : bold text (default False)
    italic       : italic text (default False)
    underline    : underlined text (default False)
    align        : 'left'|'right'|'center'|'justify' (default 'justify')
    space_before : space above in points (default 0)
    space_after  : space below in points (default 6)
    left_indent  : left indent in inches (default 0)
    line_spacing : line spacing multiplier (default 1.15)
    """
    p = doc.add_paragraph()
    p.alignment = _ALIGN_MAP.get(align.lower(), WD_ALIGN_PARAGRAPH.JUSTIFY)
    p.paragraph_format.space_before  = Pt(space_before)
    p.paragraph_format.space_after   = Pt(space_after)
    p.paragraph_format.line_spacing  = line_spacing
    p.paragraph_format.left_indent   = Inches(left_indent)

    r = p.add_run(text)
    r.font.name  = font_name
    r.font.size  = Pt(font_size)
    r.bold       = bold
    r.italic     = italic
    r.underline  = underline
    return p
