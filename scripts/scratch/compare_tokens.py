import docx
import pymupdf
import re

pdf_doc = pymupdf.open("input/DOC0000074469-Non-Parsable-en-US#FPREP_DXEGTK#.pdf")
word_doc = docx.Document("runs/DOC0000074469_20261002_1045/build/DOC0000074469-Non-Parsable-en-US#FPREP_DXEGTK#.docx")

pdf_text = ""
for page in pdf_doc:
    pdf_text += page.get_text("text") + "\n"

docx_text = []
for p in word_doc.paragraphs:
    if p.text.strip():
        docx_text.append(p.text.strip())
for t in word_doc.tables:
    for row in t.rows:
        for cell in row.cells:
            for p in cell.paragraphs:
                if p.text.strip():
                    docx_text.append(p.text.strip())

docx_full = "\n".join(docx_text)

def tokenize(text):
    return re.findall(r"\b\w+\b", text.lower())

pdf_tokens = set(tokenize(pdf_text))
docx_tokens = set(tokenize(docx_full))

missing = pdf_tokens - docx_tokens
print(f"PDF unique tokens: {len(pdf_tokens)}")
print(f"DOCX unique tokens: {len(docx_tokens)}")
print(f"Missing tokens count: {len(missing)}")
print(f"Sample missing tokens (first 50): {sorted(list(missing))[:50]}")

# Check per page
for pno in range(len(pdf_doc)):
    page_txt = pdf_doc[pno].get_text("text")
    p_tokens = set(tokenize(page_txt))
    p_missing = p_tokens - docx_tokens
    recall = 1.0 - (len(p_missing) / max(len(p_tokens), 1))
    print(f"Page {pno + 1} token recall: {recall:.2%} (missing {len(p_missing)} / {len(p_tokens)})")
