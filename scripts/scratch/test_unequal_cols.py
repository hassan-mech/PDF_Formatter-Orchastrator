import sys
from pathlib import Path
sys.path.insert(0, str(Path(__file__).resolve().parent.parent.parent))

import docx
from docx.enum.section import WD_SECTION_START
from docx.enum.text import WD_BREAK
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls, qn
from docx.shared import Inches, Pt
from scripts.safe_word_export import convert_docx_to_pdf
import pymupdf

def set_custom_columns(section, col_widths_pt: list[float], space_pt: float = 19.9):
    sectPr = section._sectPr
    for old_cols in sectPr.findall(qn('w:cols')):
        sectPr.remove(old_cols)
    w_ns = nsdecls('w')
    space_dxa = int(space_pt * 20)
    
    col_elements = ""
    for i, w_pt in enumerate(col_widths_pt):
        w_dxa = int(w_pt * 20)
        if i < len(col_widths_pt) - 1:
            col_elements += f'<w:col {w_ns} w:w="{w_dxa}" w:space="{space_dxa}"/>\n'
        else:
            col_elements += f'<w:col {w_ns} w:w="{w_dxa}"/>\n'
            
    cols_xml = f'<w:cols {w_ns} w:num="{len(col_widths_pt)}" w:equalWidth="0" w:space="{space_dxa}">\n{col_elements}</w:cols>'
    sectPr.append(parse_xml(cols_xml))

doc = docx.Document()

# Page 1:
s1 = doc.sections[0]
s1.page_width = Inches(8.27)
s1.page_height = Inches(10.63)
s1.top_margin = Inches(0.58)
s1.bottom_margin = Inches(0.40)
s1.left_margin = Inches(1.38)
s1.right_margin = Inches(0.96)

# Section 1: Unequal columns on Page 1: Left=297.6 pt (4.13 in), Right=109.3 pt (1.52 in), Gutter=19.9 pt
set_custom_columns(s1, [297.6, 109.3], space_pt=19.9)

p1 = doc.add_paragraph('Left column text with width 297.6 pt. Testing word wrapping and layout in native columns.')
p1.add_run().add_break(WD_BREAK.COLUMN)
p2 = doc.add_paragraph('Right column text with width 109.3 pt. Fits metadata.')

out_path = Path('scripts/scratch/test_unequal_cols.docx')
doc.save(str(out_path))

pdf_path = Path('scripts/scratch/test_unequal_cols.pdf')
convert_docx_to_pdf(out_path.resolve(), pdf_path.resolve(), timeout_sec=20)

pdf = pymupdf.open(str(pdf_path))
print(f'Generated PDF has {len(pdf)} pages.')
for b in pdf[0].get_text('blocks'):
    print(f'  [{b[0]:.1f}, {b[1]:.1f}, {b[2]:.1f}, {b[3]:.1f}] w={b[2]-b[0]:.1f}: {b[4].strip()}')
