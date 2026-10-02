import sys
from pathlib import Path
sys.path.insert(0, str(Path(__file__).resolve().parent.parent.parent))

import docx
from docx.enum.section import WD_SECTION_START, WD_ORIENT
from docx.enum.text import WD_BREAK, WD_ALIGN_PARAGRAPH
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls, qn
from docx.shared import Inches, Pt, RGBColor
from scripts.safe_word_export import convert_docx_to_pdf
import pymupdf

def set_section_columns(section, num_cols: int, space_pt: float = 14.4):
    sectPr = section._sectPr
    for old_cols in sectPr.findall(qn('w:cols')):
        sectPr.remove(old_cols)
    space_dxa = int(space_pt * 20)
    w_ns = nsdecls('w')
    cols_xml = f'<w:cols {w_ns} w:num="{num_cols}" w:space="{space_dxa}"/>'
    sectPr.append(parse_xml(cols_xml))

doc = docx.Document()

# Page 1:
# Section 1 (1 column): Header & Title
s1 = doc.sections[0]
s1.page_width = Inches(8.27)
s1.page_height = Inches(10.63)
s1.top_margin = Inches(0.58)
s1.bottom_margin = Inches(0.40)
s1.left_margin = Inches(1.38)
s1.right_margin = Inches(0.96)
set_section_columns(s1, 1)

p_title = doc.add_paragraph('Melanoma metastásico, un caso extraordinario')
p_title.runs[0].font.size = Pt(16)
p_title.runs[0].font.bold = True

# Section 2 (continuous, 2 columns):
s2 = doc.add_section(WD_SECTION_START.CONTINUOUS)
set_section_columns(s2, 2, space_pt=14.4)

p_c1 = doc.add_paragraph('Left column: Resumen content here...')
p_c1.add_run().add_break(WD_BREAK.COLUMN)
p_c2 = doc.add_paragraph('Right column: Affiliations and metadata here...')

# Page 2: New page section (2 columns)
s3 = doc.add_section(WD_SECTION_START.NEW_PAGE)
s3.page_width = Inches(8.27)
s3.page_height = Inches(10.63)
s3.top_margin = Inches(0.58)
s3.bottom_margin = Inches(0.40)
s3.left_margin = Inches(0.98) # Even page margin
s3.right_margin = Inches(1.34)
set_section_columns(s3, 2, space_pt=14.4)

p_p2_c1 = doc.add_paragraph('Page 2 Column 1: Abstract continuation and Antecedentes...')
p_p2_c1.add_run().add_break(WD_BREAK.COLUMN)
p_p2_c2 = doc.add_paragraph('Page 2 Column 2: Body text...')

out_path = Path('scripts/scratch/test_native_sections.docx')
doc.save(str(out_path))

pdf_path = Path('scripts/scratch/test_native_sections.pdf')
convert_docx_to_pdf(out_path.resolve(), pdf_path.resolve(), timeout_sec=20)

pdf = pymupdf.open(str(pdf_path))
print(f'Generated PDF has {len(pdf)} pages.')
for pno in range(len(pdf)):
    print(f'--- Page {pno+1} ---')
    for b in pdf[pno].get_text('blocks'):
        print(f'  [{b[0]:.1f}, {b[1]:.1f}, {b[2]:.1f}, {b[3]:.1f}]: {b[4].strip()}')
