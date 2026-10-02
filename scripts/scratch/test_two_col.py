import docx
from docx.enum.text import WD_BREAK
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls, qn
from pathlib import Path

doc = docx.Document()
section = doc.sections[0]
sectPr = section._sectPr

# Set 2 columns natively
cols = sectPr.find(qn('w:cols'))
if cols is not None:
    sectPr.remove(cols)

w_ns = nsdecls('w')
new_cols = parse_xml(f'<w:cols {w_ns} w:num="2" w:space="288"/>')
sectPr.append(new_cols)

p1 = doc.add_paragraph('Left column text paragraph 1.')
p2 = doc.add_paragraph('Left column text paragraph 2.')
p3 = doc.add_paragraph('Column break coming now:')
p3.add_run().add_break(WD_BREAK.COLUMN)
p4 = doc.add_paragraph('Right column text paragraph 1.')

out_path = Path('scripts/scratch/test_2col.docx')
doc.save(str(out_path))
print('Saved test_2col.docx successfully!')

import sys
sys.path.insert(0, str(Path(__file__).resolve().parent.parent.parent))
from scripts.safe_word_export import convert_docx_to_pdf
pdf_path = Path('scripts/scratch/test_2col.pdf')
if convert_docx_to_pdf(out_path.resolve(), pdf_path.resolve(), timeout_sec=20):
    print('Converted to PDF successfully!')
    import pymupdf
    pdf = pymupdf.open(str(pdf_path))
    print('PDF page count:', len(pdf))
    for b in pdf[0].get_text('blocks'):
        print(f'Block: [{b[0]:.1f}, {b[1]:.1f}, {b[2]:.1f}, {b[3]:.1f}]: {b[4].strip()}')
