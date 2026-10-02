"""
parallel_build.py — Parallel DOCX section builder

Processes page results from parallel_engine.py and builds DOCX sections
concurrently using a ThreadPoolExecutor, then merges them in page order.

Architecture:
    PageResults (list)
         │
         ▼  split into chunks
    ┌────────────┐    ┌────────────┐    ┌────────────┐
    │ Builder-0  │    │ Builder-1  │    │ Builder-2  │   (threads)
    │ pages 1-10 │    │ pages11-20 │    │ pages21-30 │
    └────┬───────┘    └─────┬──────┘    └─────┬──────┘
         │                  │                  │
         └──────────────────┼──────────────────┘
                            ▼
                    DocxMerger (python-docx)
                    ────────────────────────
                    final output.docx

Usage:
    from parallel_build import ParallelDocxBuilder
    from parallel_engine import ParallelPDFEngine, WorkerMode

    engine = ParallelPDFEngine("input/file.pdf", workers=8)
    page_results = engine.run(WorkerMode.EXTRACT_TEXT)

    builder = ParallelDocxBuilder(page_results, output_path="output/file.docx")
    builder.build()
"""

from __future__ import annotations

import io
import time
from concurrent.futures import ThreadPoolExecutor, as_completed
from dataclasses import dataclass
from pathlib import Path
from typing import Callable, List, Optional

try:
    import docx
    from docx.shared import Pt, Inches
    from docx.enum.text import WD_ALIGN_PARAGRAPH
except ImportError:
    raise ImportError("python-docx required: pip install python-docx")

# Optional: docxcompose for merging sub-docs
try:
    from docxcompose.composer import Composer
    HAS_COMPOSE = True
except ImportError:
    HAS_COMPOSE = False

from parallel_engine import PageResult

# Optional: generalized image-to-word parser for scanned pages
try:
    from image_to_word import parse_image_bytes
    HAS_IMAGE_PARSER = True
except ImportError:
    HAS_IMAGE_PARSER = False



# ────────────────────────────────────────────────────────────────────────────
# Section builder (runs in a thread)
# ────────────────────────────────────────────────────────────────────────────

@dataclass
class SectionSpec:
    """Describes a contiguous range of pages to build into one sub-document."""
    section_id: int
    page_results: List[PageResult]
    page_formatter: Optional[Callable]  # custom formatter fn if provided


@dataclass
class SectionDoc:
    """A built sub-document (in-memory bytes) for one section."""
    section_id: int
    docx_bytes: bytes
    page_count: int
    elapsed_ms: float
    success: bool
    error: str = ""


def _default_page_formatter(doc: "docx.Document", result: PageResult):
    """Default plain-text formatter — override with custom logic per project."""
    # Page header
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(2)
    r = p.add_run(f"— Page {result.page_number} —")
    r.bold = True
    r.font.size = Pt(9)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER

    if result.is_image_only and not result.text and not result.tables:
        if result.image_bytes:
            parsed = False
            if HAS_IMAGE_PARSER:
                try:
                    parse_image_bytes(result.image_bytes, doc_instance=doc)
                    parsed = True
                except Exception:
                    parsed = False
            if not parsed:
                p_img = doc.add_paragraph()
                p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
                run_img = p_img.add_run()
                run_img.add_picture(io.BytesIO(result.image_bytes), width=Inches(6.5))
        else:
            p2 = doc.add_paragraph()
            p2.add_run("[image-only page — OCR text not available]").italic = True
            p2.alignment = WD_ALIGN_PARAGRAPH.CENTER


    else:
        # Render text preserving paragraph breaks
        for block in result.text.split("\n\n"):
            block = block.strip()
            if not block:
                continue
            p2 = doc.add_paragraph()
            p2.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            p2.paragraph_format.space_after = Pt(4)
            p2.paragraph_format.line_spacing = 1.15
            p2.add_run(block)

    # Render detected tables
    if result.tables:
        for tbl_data in result.tables:
            if not tbl_data or not tbl_data[0]:
                continue
            num_rows = len(tbl_data)
            num_cols = max(len(row) for row in tbl_data)
            table = doc.add_table(rows=num_rows, cols=num_cols)
            table.autofit = True
            for r_idx, row in enumerate(tbl_data):
                for c_idx, cell_text in enumerate(row):
                    if c_idx < num_cols:
                        cell = table.cell(r_idx, c_idx)
                        cell.text = str(cell_text or "").strip()
            doc.add_paragraph()

    # Page separator
    doc.add_paragraph()


def _build_section(spec: SectionSpec, page_formatter: Callable) -> SectionDoc:
    """Build a sub-DOCX for a range of pages in a thread."""
    t0 = time.perf_counter()
    try:
        doc = docx.Document()
        # Minimal section setup — A4 portrait
        for section in doc.sections:
            section.page_width = Inches(8.27)
            section.page_height = Inches(11.69)
            section.top_margin = Inches(1.0)
            section.bottom_margin = Inches(1.0)
            section.left_margin = Inches(1.0)
            section.right_margin = Inches(1.0)

        # Remove default empty paragraph
        if doc.paragraphs:
            p = doc.paragraphs[0]._element
            p.getparent().remove(p)

        for result in spec.page_results:
            page_formatter(doc, result)

        buf = io.BytesIO()
        doc.save(buf)
        docx_bytes = buf.getvalue()

        return SectionDoc(
            section_id=spec.section_id,
            docx_bytes=docx_bytes,
            page_count=len(spec.page_results),
            elapsed_ms=(time.perf_counter() - t0) * 1000,
            success=True,
        )
    except Exception as exc:
        return SectionDoc(
            section_id=spec.section_id,
            docx_bytes=b"",
            page_count=0,
            elapsed_ms=(time.perf_counter() - t0) * 1000,
            success=False,
            error=str(exc),
        )


# ────────────────────────────────────────────────────────────────────────────
# DOCX Merger
# ────────────────────────────────────────────────────────────────────────────

class DocxMerger:
    """Merges multiple in-memory DOCX bytes into one final document."""

    @staticmethod
    def merge(section_docs: List[SectionDoc], output_path: Path):
        """Merge section docs in order, save to output_path."""
        if HAS_COMPOSE:
            DocxMerger._merge_with_compose(section_docs, output_path)
        else:
            DocxMerger._merge_fallback(section_docs, output_path)

    @staticmethod
    def _merge_with_compose(section_docs: List[SectionDoc], output_path: Path):
        """Use docxcompose for proper section merging."""
        sorted_docs = sorted(section_docs, key=lambda d: d.section_id)
        # Load first doc as base
        base = docx.Document(io.BytesIO(sorted_docs[0].docx_bytes))
        composer = Composer(base)
        for sd in sorted_docs[1:]:
            sub = docx.Document(io.BytesIO(sd.docx_bytes))
            composer.append(sub)
        composer.save(str(output_path))

    @staticmethod
    def _merge_fallback(section_docs: List[SectionDoc], output_path: Path):
        """Fallback: copy XML elements from each sub-doc into a master doc."""
        from docx.oxml.ns import qn

        sorted_docs = sorted(section_docs, key=lambda d: d.section_id)
        master = docx.Document(io.BytesIO(sorted_docs[0].docx_bytes))
        master_body = master.element.body

        for sd in sorted_docs[1:]:
            sub = docx.Document(io.BytesIO(sd.docx_bytes))
            sub_body = sub.element.body
            for child in list(sub_body):
                # Skip the final sectPr (section properties) — keep master's
                if child.tag == qn("w:sectPr"):
                    continue
                master_body.append(child)

        output_path.parent.mkdir(parents=True, exist_ok=True)
        master.save(str(output_path))


# ────────────────────────────────────────────────────────────────────────────
# High-level Parallel Builder API
# ────────────────────────────────────────────────────────────────────────────

class ParallelDocxBuilder:
    """
    Parallel DOCX builder.

    Splits PageResults into chunks, builds sub-documents in parallel threads,
    then merges them into a single output DOCX.

    Parameters
    ----------
    page_results    : ordered list from ParallelPDFEngine.run()
    output_path     : final .docx path
    workers         : number of builder threads (default: 4)
    chunk_size      : pages per section/sub-doc (default: 20)
    page_formatter  : callable(doc, PageResult) — custom formatter
                      If None, uses the default plain-text formatter.

    Example
    -------
    builder = ParallelDocxBuilder(
        page_results, "output/result.docx",
        workers=4, chunk_size=20,
        page_formatter=my_custom_formatter,
    )
    builder.build()
    """

    def __init__(
        self,
        page_results: List[PageResult],
        output_path: str | Path,
        workers: int = 4,
        chunk_size: int = 20,
        page_formatter: Optional[Callable] = None,
    ):
        self.page_results = page_results
        self.output_path = Path(output_path)
        self.workers = workers
        self.chunk_size = chunk_size
        self.page_formatter = page_formatter or _default_page_formatter

    def _make_specs(self) -> List[SectionSpec]:
        specs = []
        for i in range(0, len(self.page_results), self.chunk_size):
            chunk = self.page_results[i : i + self.chunk_size]
            specs.append(SectionSpec(
                section_id=i // self.chunk_size,
                page_results=chunk,
                page_formatter=self.page_formatter,
            ))
        return specs

    def build(self, verbose: bool = True) -> Path:
        specs = self._make_specs()
        section_docs: List[Optional[SectionDoc]] = [None] * len(specs)
        t_start = time.perf_counter()

        if verbose:
            print(f"\n📝 ParallelDocxBuilder starting")
            print(f"   Output     : {self.output_path}")
            print(f"   Pages      : {len(self.page_results)}")
            print(f"   Sections   : {len(specs)}")
            print(f"   Threads    : {self.workers}")
            print(f"   Chunk size : {self.chunk_size}")
            has_compose = "yes (docxcompose)" if HAS_COMPOSE else "no (fallback XML merge)"
            print(f"   Composer   : {has_compose}")

        with ThreadPoolExecutor(max_workers=self.workers) as pool:
            futures = {
                pool.submit(_build_section, spec, self.page_formatter): spec.section_id
                for spec in specs
            }
            for future in as_completed(futures):
                sid = futures[future]
                sd = future.result()
                section_docs[sid] = sd
                if verbose:
                    status = "✅" if sd.success else "❌"
                    print(
                        f"  {status} Section {sd.section_id+1:>3}/{len(specs)}"
                        f"  pages={sd.page_count}  {sd.elapsed_ms:.0f}ms"
                        + (f"  ERR: {sd.error}" if sd.error else ""),
                        flush=True,
                    )

        # Filter successful sections only
        valid_docs = [d for d in section_docs if d and d.success]
        if not valid_docs:
            raise RuntimeError("All sections failed — cannot produce output DOCX.")

        if verbose:
            print(f"\n🔗 Merging {len(valid_docs)} sections...")

        DocxMerger.merge(valid_docs, self.output_path)

        elapsed = time.perf_counter() - t_start
        if verbose:
            print(f"✅ DOCX saved: {self.output_path}  ({elapsed:.2f}s total)")

        return self.output_path


# ────────────────────────────────────────────────────────────────────────────
# CLI
# ────────────────────────────────────────────────────────────────────────────

def main():
    """
    Demo: run full pipeline (inspect → parallel build) end-to-end.
    For production use, import ParallelDocxBuilder directly.
    """
    import argparse
    from parallel_engine import ParallelPDFEngine, WorkerMode

    parser = argparse.ArgumentParser(
        description="parallel_build.py — parallel DOCX builder demo",
    )
    parser.add_argument("-i", "--input", required=True, help="PDF path")
    parser.add_argument("-o", "--output", default=None, help="Output .docx path")
    parser.add_argument("--workers", type=int, default=4, help="Builder threads (default: 4)")
    parser.add_argument("--chunk-size", type=int, default=20, help="Pages per section (default: 20)")
    parser.add_argument("--engine-workers", type=int, default=None, help="Extraction workers (default: cpu-2)")
    parser.add_argument("--dpi", type=int, default=150)
    parser.add_argument("--pages", default=None)
    args = parser.parse_args()

    pdf_path = Path(args.input)
    output = Path(args.output) if args.output else Path("output") / (pdf_path.stem + ".docx")

    # Stage 1: parallel extraction
    engine = ParallelPDFEngine(
        pdf_path=pdf_path,
        workers=args.engine_workers,
        chunk_size=8,
        dpi=args.dpi,
        pages=args.pages,
    )
    page_results = engine.run(mode=WorkerMode.EXTRACT_TEXT, verbose=True)

    # Stage 2: parallel build
    builder = ParallelDocxBuilder(
        page_results=page_results,
        output_path=output,
        workers=args.workers,
        chunk_size=args.chunk_size,
    )
    builder.build(verbose=True)


if __name__ == "__main__":
    main()
