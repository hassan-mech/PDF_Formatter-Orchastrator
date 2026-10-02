"""
build_ug100_doc.py — Universal builder for UG-100 translation documents (FPREP workflow).

Handles:
  1. UG-100 Form Documents (e.g. Local AE reporting form DOC-000226884):
     - Greek Adverse Event Reporting Form (Page 1): Created as native Word text & tables
       with exact typography, borders, and calculated dash lengths filling the margins.
     - English SOP / Revision History (Page 2): Embedded as exact high-resolution screenshot (300 DPI).
  2. UG-100 Email Thread Documents (e.g. Outlook email threads):
     - Print banner & Greek source email: Created as editable Word text with exact typography and styling.
     - Forwarded English thread: Embedded as high-resolution screenshots.

Usage:
    python scripts/build_ug100_doc.py -i input/file.pdf -o output/file.docx
"""

from __future__ import annotations

import argparse
import sys
from pathlib import Path
from typing import Optional

try:
    import fitz  # PyMuPDF
except ImportError:
    print("[ERROR] PyMuPDF required: pip install pymupdf")
    sys.exit(1)

try:
    import docx
    from docx.enum.table import WD_TABLE_ALIGNMENT
    from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
    from docx.oxml import parse_xml
    from docx.oxml.ns import nsdecls, qn
    from docx.shared import Inches, Mm, Pt, RGBColor
except ImportError:
    print("[ERROR] python-docx required: pip install python-docx")
    sys.exit(1)


def set_cell_margins(cell, top=40, bottom=40, left=80, right=80):
    """Set cell padding in dxa (1 pt = 20 dxa)."""
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(
        f'<w:tcMar {nsdecls("w")}>'
        f'<w:top w:w="{top}" w:type="dxa"/>'
        f'<w:bottom w:w="{bottom}" w:type="dxa"/>'
        f'<w:left w:w="{left}" w:type="dxa"/>'
        f'<w:right w:w="{right}" w:type="dxa"/>'
        f'</w:tcMar>'
    )
    tcPr.append(tcMar)


def set_cell_borders(cell, top="none", bottom="none", left="none", right="none", sz="4", color="000000"):
    """Set specific borders for a table cell."""
    tcPr = cell._tc.get_or_add_tcPr()
    tcBorders = parse_xml(
        f'<w:tcBorders {nsdecls("w")}>'
        f'<w:top w:val="{top}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'<w:left w:val="{left}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'<w:bottom w:val="{bottom}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'<w:right w:val="{right}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'</w:tcBorders>'
    )
    tcPr.append(tcBorders)


def set_font_xml(run, font_name: str):
    rPr = run._r.get_or_add_rPr()
    rFonts = parse_xml(
        f'<w:rFonts {nsdecls("w")} w:ascii="{font_name}" w:hAnsi="{font_name}" w:cs="{font_name}"/>'
    )
    rPr.append(rFonts)


# =============================================================================
# 1. UG-100 FORM BUILDER (Local AE Reporting Form)
# =============================================================================

def build_ug100_form(
    pdf_path: Path,
    output_docx: Path,
    renders_dir: Path,
    margin_mm: float = 15.0,
    dpi: int = 300,
) -> Path:
    print(f"📄 Building UG-100 Form document: {pdf_path.name}")
    doc_pdf = fitz.open(str(pdf_path))

    # Extract CSL Logo from Page 1
    logo_path = renders_dir / "csl_logo.png"
    for img_info in doc_pdf[0].get_images():
        xref = img_info[0]
        base_img = doc_pdf.extract_image(xref)
        with open(logo_path, "wb") as f:
            f.write(base_img["image"])
        break

    # Extract Page 2 English SOP screenshot
    page2 = doc_pdf[1]
    pix2 = page2.get_pixmap(dpi=dpi)
    p2_img_path = renders_dir / f"{pdf_path.stem}_p2_english.png"
    pix2.save(str(p2_img_path))
    doc_pdf.close()

    # Build Word Document
    doc = docx.Document()
    sec1 = doc.sections[0]

    sec1.page_width = Inches(8.5)
    sec1.page_height = Inches(11.0)
    sec1.top_margin = Mm(margin_mm)
    sec1.bottom_margin = Mm(margin_mm)
    sec1.left_margin = Mm(margin_mm)
    sec1.right_margin = Mm(margin_mm)

    printable_width_in = 8.5 - 2.0 * (margin_mm / 25.4)

    FONT_ARIAL = "Arial"
    FONT_BOOKMAN = "Bookman Old Style"
    FONT_CALIBRI = "Calibri"
    FONT_SYMBOL = "Segoe UI Symbol"
    BLACK = RGBColor(0, 0, 0)
    BLUE_COLOR = RGBColor(51, 51, 255)

    # 1. Header Table (2 columns: left logo & subtitle, right metadata)
    tbl_hdr = doc.add_table(rows=1, cols=2)
    tbl_hdr.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl_hdr.autofit = False

    w_left = Inches(printable_width_in * 0.55)
    w_right = Inches(printable_width_in * 0.45)
    tbl_hdr.columns[0].width = w_left
    tbl_hdr.columns[1].width = w_right

    c_left = tbl_hdr.cell(0, 0)
    c_right = tbl_hdr.cell(0, 1)
    c_left.width = w_left
    c_right.width = w_right

    # Left: Logo + "Local AE reporting form"
    p_logo = c_left.paragraphs[0]
    p_logo.paragraph_format.space_before = Pt(0)
    p_logo.paragraph_format.space_after = Pt(6)
    r_logo = p_logo.add_run()
    r_logo.add_picture(str(logo_path), width=Inches(0.68))

    p_sub = c_left.add_paragraph()
    p_sub.paragraph_format.space_before = Pt(0)
    p_sub.paragraph_format.space_after = Pt(0)
    r_sub = p_sub.add_run("Local AE reporting form")
    r_sub.font.name = FONT_ARIAL
    r_sub.font.size = Pt(13.2)
    r_sub.font.color.rgb = BLACK
    set_font_xml(r_sub, FONT_ARIAL)

    # Right: Form metadata
    p_meta = c_right.paragraphs[0]
    p_meta.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    p_meta.paragraph_format.space_before = Pt(0)
    p_meta.paragraph_format.space_after = Pt(0)
    p_meta.paragraph_format.line_spacing = Pt(12)

    lines_meta = [
        ("Form", True, 11.3),
        ("DOC-000226884", False, 10.4),
        ("v2.0", False, 10.4),
        ("Effective Date (GMT): 18 Feb 2024", False, 10.4),
    ]
    for idx, (txt, is_bold, sz) in enumerate(lines_meta):
        if idx > 0:
            p_m = c_right.add_paragraph()
            p_m.alignment = WD_ALIGN_PARAGRAPH.RIGHT
            p_m.paragraph_format.space_before = Pt(0)
            p_m.paragraph_format.space_after = Pt(0)
            p_m.paragraph_format.line_spacing = Pt(12)
        else:
            p_m = p_meta
        r_m = p_m.add_run(txt)
        r_m.font.name = FONT_ARIAL
        r_m.font.size = Pt(sz)
        r_m.font.bold = is_bold
        r_m.font.color.rgb = BLACK
        set_font_xml(r_m, FONT_ARIAL)

    for cell in (c_left, c_right):
        set_cell_borders(cell, "none", "none", "none", "none")
        set_cell_margins(cell, top=0, bottom=0, left=0, right=0)

    # Solid black divider line below header (0.75 pt)
    p_div = doc.add_paragraph()
    p_div.paragraph_format.space_before = Pt(2)
    p_div.paragraph_format.space_after = Pt(6)
    pPr = p_div._p.get_or_add_pPr()
    pBdr = parse_xml(f'<w:pBdr {nsdecls("w")}><w:bottom w:val="single" w:sz="6" w:space="1" w:color="000000"/></w:pBdr>')
    pPr.append(pBdr)

    # 2. Greek Document Title
    p_title = doc.add_paragraph()
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_title.paragraph_format.space_before = Pt(2)
    p_title.paragraph_format.space_after = Pt(6)
    p_title.paragraph_format.line_spacing = Pt(12)

    title_lines = [
        "ΑΝΑΦΟΡΑ ΑΝΕΠΙΘΥΜΗΤΗΣ ΕΝΕΡΓΕΙΑΣ",
        "ΠΡΟΣ ΤΟ ΤΜΗΜΑ ΦΑΡΜΑΚΟΕΠΑΓΡΥΠΝΗΣΗΣ ΤΗΣ",
        "CSL BEHRING ΕΠΕ",
    ]
    for idx, tl in enumerate(title_lines):
        if idx > 0:
            p_title.add_run("\n")
        r_t = p_title.add_run(tl)
        r_t.font.name = FONT_BOOKMAN
        r_t.font.size = Pt(10.4)
        r_t.font.bold = True
        r_t.font.color.rgb = BLACK
        set_font_xml(r_t, FONT_BOOKMAN)

    # 3. Table 1: Reporter Box (2 rows)
    tbl1 = doc.add_table(rows=2, cols=2)
    tbl1.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl1.autofit = False

    w_t1_left = Inches(printable_width_in * 0.58)
    w_t1_right = Inches(printable_width_in * 0.42)
    tbl1.columns[0].width = w_t1_left
    tbl1.columns[1].width = w_t1_right

    c_r0 = tbl1.cell(0, 0)
    c_r0.merge(tbl1.cell(0, 1))
    c_r0.width = Inches(printable_width_in)

    p_t1_r0 = c_r0.paragraphs[0]
    p_t1_r0.paragraph_format.space_before = Pt(0)
    p_t1_r0.paragraph_format.space_after = Pt(0)
    p_t1_r0.paragraph_format.line_spacing = Pt(11.5)

    r_t1_1 = p_t1_r0.add_run("Ονοματεπώνυμο Αναφέροντος Υπαλλήλου:Eva Vairaktari ")
    r_t1_1.font.name = FONT_BOOKMAN
    r_t1_1.font.size = Pt(9.3)
    r_t1_1.font.bold = True
    r_t1_1.font.color.rgb = BLACK
    set_font_xml(r_t1_1, FONT_BOOKMAN)

    r_t1_d = p_t1_r0.add_run("-" * 66)
    r_t1_d.font.name = FONT_BOOKMAN
    r_t1_d.font.size = Pt(9.3)
    r_t1_d.font.bold = False
    r_t1_d.font.color.rgb = BLACK
    set_font_xml(r_t1_d, FONT_BOOKMAN)

    # Row 1 Left: Ημερομηνία αναφοράς
    c_r1_left = tbl1.cell(1, 0)
    c_r1_left.width = w_t1_left
    p_t1_r1l = c_r1_left.paragraphs[0]
    p_t1_r1l.paragraph_format.space_before = Pt(0)
    p_t1_r1l.paragraph_format.space_after = Pt(0)
    p_t1_r1l.paragraph_format.line_spacing = Pt(11.5)

    r_date_lbl = p_t1_r1l.add_run("Ημερομηνία αναφοράς:")
    r_date_lbl.font.name = FONT_CALIBRI
    r_date_lbl.font.size = Pt(9.3)
    r_date_lbl.font.bold = True
    r_date_lbl.font.color.rgb = BLACK
    set_font_xml(r_date_lbl, FONT_CALIBRI)

    r_date_val = p_t1_r1l.add_run("30/9/26 " + "-" * 54)
    r_date_val.font.name = FONT_CALIBRI
    r_date_val.font.size = Pt(9.3)
    r_date_val.font.bold = False
    r_date_val.font.color.rgb = BLACK
    set_font_xml(r_date_val, FONT_CALIBRI)

    # Row 1 Right: Υπογραφή
    c_r1_right = tbl1.cell(1, 1)
    c_r1_right.width = w_t1_right
    p_t1_r1r = c_r1_right.paragraphs[0]
    p_t1_r1r.paragraph_format.space_before = Pt(0)
    p_t1_r1r.paragraph_format.space_after = Pt(0)
    p_t1_r1r.paragraph_format.line_spacing = Pt(11.5)

    r_sig_lbl = p_t1_r1r.add_run("Υπογραφή: ")
    r_sig_lbl.font.name = FONT_CALIBRI
    r_sig_lbl.font.size = Pt(9.3)
    r_sig_lbl.font.bold = True
    r_sig_lbl.font.color.rgb = BLACK
    set_font_xml(r_sig_lbl, FONT_CALIBRI)

    r_sig_val = p_t1_r1r.add_run("E.Vairaktari " + "-" * 40)
    r_sig_val.font.name = FONT_CALIBRI
    r_sig_val.font.size = Pt(9.3)
    r_sig_val.font.bold = False
    r_sig_val.font.color.rgb = BLACK
    set_font_xml(r_sig_val, FONT_CALIBRI)

    set_cell_borders(c_r0, top="single", bottom="none", left="single", right="single", sz="4")
    set_cell_margins(c_r0, top=50, bottom=30, left=80, right=80)
    set_cell_borders(c_r1_left, top="none", bottom="single", left="single", right="none", sz="4")
    set_cell_margins(c_r1_left, top=30, bottom=50, left=80, right=40)
    set_cell_borders(c_r1_right, top="none", bottom="single", left="none", right="single", sz="4")
    set_cell_margins(c_r1_right, top=30, bottom=50, left=40, right=80)

    p_sp = doc.add_paragraph()
    p_sp.paragraph_format.space_before = Pt(3)
    p_sp.paragraph_format.space_after = Pt(0)
    p_sp.paragraph_format.line_spacing = Pt(3)

    # 4. Table 2: Detailed Reporting Form Box
    tbl2 = doc.add_table(rows=1, cols=1)
    tbl2.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl2.autofit = False
    tbl2.columns[0].width = Inches(printable_width_in)

    c_box = tbl2.cell(0, 0)
    c_box.width = Inches(printable_width_in)
    set_cell_borders(c_box, top="single", bottom="single", left="single", right="single", sz="4")
    set_cell_margins(c_box, top=50, bottom=50, left=80, right=80)

    table2_content = [
        # Section 1
        ("Στοιχεία αναφέροντος Επαγγελματία Υγείας/ άλλου Αναφέροντος", True, FONT_CALIBRI, Pt(0), Pt(0)),
        ("Όνομα:Dr Angelopoulou" + "-" * 144, False, FONT_CALIBRI, Pt(0), Pt(0)),
        ("Διεύθυνση:Alexandroupolis Hospital " + "-" * 128, False, FONT_CALIBRI, Pt(0), Pt(0)),
        ("Τηλέφωνο:6987050481" + "-" * 145, False, FONT_CALIBRI, Pt(0), Pt(0)),
        ("Fax: " + "-" * 170, False, FONT_CALIBRI, Pt(0), Pt(0)),
        ("e-mail: " + "-" * 166, False, FONT_CALIBRI, Pt(0), Pt(0)),
        # Section 2
        ("Στοιχεία ασθενούς", True, FONT_CALIBRI, Pt(2), Pt(0)),
        ("Αρχικά: " + "-" * 165, False, FONT_CALIBRI, Pt(0), Pt(0)),
        ("Ημερομηνία γέννησης/ ηλικία:50yo " + "-" * 128, False, FONT_CALIBRI, Pt(0), Pt(0)),
        ("Φύλο:M " + "-" * 164, False, FONT_CALIBRI, Pt(0), Pt(0)),
        ("Άλλες πληροφορίες: The HCP did not share more information " + "-" * 95, False, FONT_CALIBRI, Pt(0), Pt(0)),
        ("-" * 175, False, FONT_CALIBRI, Pt(0), Pt(0)),
        # Section 3
        ("Στοιχεία προϊόντος", True, FONT_CALIBRI, Pt(2), Pt(0)),
        ("Εμπορική ονομασία:Hizentra " + "-" * 137, False, FONT_CALIBRI, Pt(0), Pt(0)),
        ("Αριθμός παρτίδας: The HCP did not share this information " + "-" * 99, False, FONT_CALIBRI, Pt(0), Pt(0)),
        ("Ένδειξη χορήγησης:CIDP " + "-" * 143, False, FONT_CALIBRI, Pt(0), Pt(0)),
        ("Ημερομηνία(ες) χορήγησης: The HCP did not share this information " + "-" * 87, False, FONT_CALIBRI, Pt(0), Pt(0)),
        ("Οδός χορήγησης:SC " + "-" * 149, False, FONT_CALIBRI, Pt(0), Pt(0)),
        ("Δοσολογικό σχήμα: " + "-" * 149, False, FONT_CALIBRI, Pt(0), Pt(0)),
        ("SPECIAL_CHECKBOX", False, FONT_CALIBRI, Pt(0), Pt(0)),
        ("Λοιπές πληροφορίες:" + "-" * 148, False, FONT_CALIBRI, Pt(0), Pt(0)),
        ("-" * 175, False, FONT_CALIBRI, Pt(0), Pt(0)),
        ("-" * 175, False, FONT_CALIBRI, Pt(0), Pt(0)),
        # Section 4
        ("Περιγραφή του συμβάντος", True, FONT_CALIBRI, Pt(2), Pt(0)),
        ("Λεπτομερής περιγραφή: Ο ασθενής έπαιρνε IVIg για 2 χρόνια και τους τελευταίους 3 μήνες Hizentra. Όμως", False, FONT_CALIBRI, Pt(0), Pt(0)),
        ("διαμαρτυρήθηκε για γρουμπαλάκια κατά την έγχυση και επέλεξε να γυρίσει στο IVIg", False, FONT_CALIBRI, Pt(0), Pt(0)),
        ("-" * 175, False, FONT_CALIBRI, Pt(0), Pt(0)),
        ("-" * 175, False, FONT_CALIBRI, Pt(0), Pt(0)),
        ("-" * 175, False, FONT_CALIBRI, Pt(0), Pt(0)),
        ("Ημερομηνία έναρξης: The HCP did not share this information " + "-" * 95, False, FONT_CALIBRI, Pt(0), Pt(0)),
        ("Διάρκεια: " + "-" * 162, False, FONT_CALIBRI, Pt(0), Pt(0)),
        ("Έκβαση: " + "-" * 164, False, FONT_CALIBRI, Pt(0), Pt(0)),
    ]

    for idx, (txt, is_bold, f_name, sp_before, sp_after) in enumerate(table2_content):
        p_line = c_box.paragraphs[0] if idx == 0 else c_box.add_paragraph()
        p_line.paragraph_format.space_before = sp_before
        p_line.paragraph_format.space_after = sp_after
        p_line.paragraph_format.line_spacing = Pt(11.2)

        if txt == "SPECIAL_CHECKBOX":
            r1 = p_line.add_run("Φύλαξη στη σωστή θερμοκρασία:  ΝΑΙ")
            r1.font.name = FONT_CALIBRI
            r1.font.size = Pt(9.3)
            r1.font.bold = False
            r1.font.color.rgb = BLACK
            set_font_xml(r1, FONT_CALIBRI)

            r_x = p_line.add_run("x")
            r_x.font.name = FONT_SYMBOL
            r_x.font.size = Pt(9.3)
            r_x.font.bold = True
            r_x.font.color.rgb = BLACK
            set_font_xml(r_x, FONT_SYMBOL)

            r2 = p_line.add_run("         ΟΧΙ  ")
            r2.font.name = FONT_CALIBRI
            r2.font.size = Pt(9.3)
            r2.font.bold = False
            r2.font.color.rgb = BLACK
            set_font_xml(r2, FONT_CALIBRI)

            r_box = p_line.add_run("☐")
            r_box.font.name = FONT_SYMBOL
            r_box.font.size = Pt(9.3)
            r_box.font.bold = False
            r_box.font.color.rgb = BLACK
            set_font_xml(r_box, FONT_SYMBOL)
        else:
            r_txt = p_line.add_run(txt)
            r_txt.font.name = f_name
            r_txt.font.size = Pt(9.3)
            r_txt.font.bold = is_bold
            r_txt.font.color.rgb = BLACK
            set_font_xml(r_txt, f_name)

    # 5. Blue Contact Box (Below Table 2)
    p_blue = doc.add_paragraph()
    p_blue.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_blue.paragraph_format.space_before = Pt(6)
    p_blue.paragraph_format.space_after = Pt(4)
    p_blue.paragraph_format.line_spacing = Pt(13)

    blue_lines = [
        "Στέλνετε όλες τις Ανεπιθύμητες Ενέργειες προς τον",
        "Τοπικό Υπεύθυνο Φαρμακοεπαγρύπνησης",
        "Τηλ: 210 7255660, 210 6527444",
        "Φαξ: 210 7255663, 2106512210",
        "E-mail: pharmacovigilance.greece@cslbehring.com",
    ]
    for idx, bl in enumerate(blue_lines):
        if idx > 0:
            p_blue.add_run("\n")
        r_bl = p_blue.add_run(bl)
        r_bl.font.name = FONT_CALIBRI
        r_bl.font.size = Pt(10.4)
        r_bl.font.bold = True
        r_bl.font.color.rgb = BLUE_COLOR
        set_font_xml(r_bl, FONT_CALIBRI)

    # 6. Page 1 Footer (3-column table with top border)
    footer1 = sec1.footer
    tbl_ftr = footer1.add_table(rows=1, cols=3, width=Inches(printable_width_in))
    tbl_ftr.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl_ftr.autofit = False

    w_f0 = Inches(printable_width_in * 0.28)
    w_f1 = Inches(printable_width_in * 0.58)
    w_f2 = Inches(printable_width_in * 0.14)
    tbl_ftr.columns[0].width = w_f0
    tbl_ftr.columns[1].width = w_f1
    tbl_ftr.columns[2].width = w_f2

    c_f0 = tbl_ftr.cell(0, 0)
    c_f1 = tbl_ftr.cell(0, 1)
    c_f2 = tbl_ftr.cell(0, 2)
    c_f0.width = w_f0
    c_f1.width = w_f1
    c_f2.width = w_f2

    for c in (c_f0, c_f1, c_f2):
        set_cell_borders(c, top="single", bottom="none", left="none", right="none", sz="4")
        set_cell_margins(c, top=40, bottom=0, left=0, right=0)

    p_f0 = c_f0.paragraphs[0]
    p_f0.paragraph_format.space_before = Pt(0)
    p_f0.paragraph_format.space_after = Pt(0)
    r_f0 = p_f0.add_run("CSLB_GLB_FORM_Template_English_V1.0")
    r_f0.font.name = FONT_ARIAL
    r_f0.font.size = Pt(4.7)
    r_f0.font.color.rgb = BLACK
    set_font_xml(r_f0, FONT_ARIAL)

    p_f1 = c_f1.paragraphs[0]
    p_f1.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_f1.paragraph_format.space_before = Pt(0)
    p_f1.paragraph_format.space_after = Pt(0)
    r_f1 = p_f1.add_run("Confidential document for authorized use only. / Vertrauliches Dokument nur zur autorisierten Verwendung.")
    r_f1.font.name = FONT_ARIAL
    r_f1.font.size = Pt(4.7)
    r_f1.font.color.rgb = BLACK
    set_font_xml(r_f1, FONT_ARIAL)

    p_f2 = c_f2.paragraphs[0]
    p_f2.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    p_f2.paragraph_format.space_before = Pt(0)
    p_f2.paragraph_format.space_after = Pt(0)
    r_f2 = p_f2.add_run("Page 1 / 1")
    r_f2.font.name = FONT_ARIAL
    r_f2.font.size = Pt(9.3)
    r_f2.font.color.rgb = BLACK
    set_font_xml(r_f2, FONT_ARIAL)

    if len(footer1.paragraphs) > 0 and not footer1.paragraphs[0].text:
        p_lead = footer1.paragraphs[0]._p
        p_lead.getparent().remove(p_lead)

    # =========================================================================
    # PAGE 2: Exact High-Resolution Screenshot of English SOP
    # =========================================================================
    sec2 = doc.add_section()
    sec2.page_width = Inches(8.5)
    sec2.page_height = Inches(11.0)
    sec2.top_margin = Mm(margin_mm)
    sec2.bottom_margin = Mm(margin_mm)
    sec2.left_margin = Mm(margin_mm)
    sec2.right_margin = Mm(margin_mm)

    sec2.header.is_linked_to_previous = False
    sec2.footer.is_linked_to_previous = False

    for p in sec2.footer.paragraphs:
        p.text = ""
        pPr = p._p.find(qn("w:pPr"))
        if pPr is not None:
            pBdr = pPr.find(qn("w:pBdr"))
            if pBdr is not None:
                pPr.remove(pBdr)

    p_p2 = doc.add_paragraph()
    p_p2.paragraph_format.space_before = Pt(0)
    p_p2.paragraph_format.space_after = Pt(0)
    p_p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_p2 = p_p2.add_run()
    # 7.28 in preserves aspect ratio without pushing image off-page
    r_p2.add_picture(str(p2_img_path), width=Inches(7.28))

    doc.save(str(output_docx))
    print(f"✅ Document successfully created: {output_docx} ({output_docx.stat().st_size:,} bytes)")
    return output_docx


# =============================================================================
# 2. UG-100 EMAIL THREAD BUILDER (Forwarded thread screenshots + Greek body)
# =============================================================================

def build_ug100_email(
    pdf_path: Path,
    output_docx: Path,
    renders_dir: Path,
    margin_mm: float = 15.0,
    dpi: int = 300,
) -> Path:
    print(f"📄 Building UG-100 Email document: {pdf_path.name}")
    doc_pdf = fitz.open(str(pdf_path))

    # Crop 1: Page 1 English thread (y=224.0 to y=742.0)
    clip1 = fitz.Rect(36.0, 224.0, 576.0, 742.0)
    pix1 = doc_pdf[0].get_pixmap(clip=clip1, dpi=dpi)
    crop1_path = renders_dir / f"{pdf_path.stem}_crop_p1.png"
    pix1.save(str(crop1_path))

    # Crop 2: Page 2 English thread (y=30.0 to y=445.0)
    clip2 = fitz.Rect(36.0, 30.0, 576.0, 445.0)
    pix2 = doc_pdf[1].get_pixmap(clip=clip2, dpi=dpi)
    crop2_path = renders_dir / f"{pdf_path.stem}_crop_p2.png"
    pix2.save(str(crop2_path))

    doc_pdf.close()

    doc = docx.Document()
    section = doc.sections[0]

    section.page_width = Inches(8.5)
    section.page_height = Inches(11.0)
    section.top_margin = Mm(margin_mm)
    section.bottom_margin = Mm(margin_mm)
    section.left_margin = Mm(margin_mm)
    section.right_margin = Mm(margin_mm)

    printable_width_in = 8.5 - 2.0 * (margin_mm / 25.4)

    # Footer page numbering
    footer = section.footer
    p_ftr = footer.paragraphs[0]
    p_ftr.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_ftr.paragraph_format.space_before = Pt(0)
    p_ftr.paragraph_format.space_after = Pt(0)
    r_ftr = p_ftr.add_run()
    r_ftr.font.name = "Segoe UI"
    r_ftr.font.size = Pt(8.5)
    r_ftr.font.color.rgb = RGBColor(60, 60, 60)
    fldSimple = parse_xml(f'<w:fldSimple {nsdecls("w")} w:instr="PAGE"/>')
    p_ftr._p.append(fldSimple)

    FONT_SEGOE = "Segoe UI"
    FONT_APTOS = "Aptos"
    BLACK = RGBColor(0, 0, 0)
    BLUE_LINK = RGBColor(0, 0, 255)

    # Print banner
    p_banner = doc.add_paragraph()
    p_banner.paragraph_format.space_before = Pt(0)
    p_banner.paragraph_format.space_after = Pt(4)
    r_banner = p_banner.add_run("Bharate, Vedangi IN/GLB EXT")
    r_banner.font.name = FONT_SEGOE
    r_banner.font.size = Pt(12)
    r_banner.font.bold = True
    r_banner.font.color.rgb = BLACK
    set_font_xml(r_banner, FONT_SEGOE)

    pPr = p_banner._p.get_or_add_pPr()
    pBdr = parse_xml(f'<w:pBdr {nsdecls("w")}><w:bottom w:val="single" w:sz="10" w:space="3" w:color="000000"/></w:pBdr>')
    pPr.append(pBdr)

    # Email header fields
    headers_data = [
        ("From:", "Vairaktari, Eva GR/ATH"),
        ("Sent:", "Wednesday, September 30, 2026 7:17 PM"),
        ("To:", "Kotsaridi, Elina GR/GLB EXT"),
        ("Cc:", "DL GRP EU ATH BPL Pharmacovigilance"),
        ("Subject:", "Re: AE reporting - Hizentra"),
    ]

    tab_pos = Inches(2.12)
    for idx, (label, val) in enumerate(headers_data):
        p_hdr = doc.add_paragraph()
        p_hdr.paragraph_format.space_before = Pt(0)
        p_hdr.paragraph_format.space_after = Pt(0.5)
        p_hdr.paragraph_format.line_spacing = Pt(13.2)
        p_hdr.paragraph_format.tab_stops.add_tab_stop(tab_pos, WD_TAB_ALIGNMENT.LEFT)

        r_lbl = p_hdr.add_run(label)
        r_lbl.font.name = FONT_SEGOE
        r_lbl.font.size = Pt(9.9)
        r_lbl.font.bold = True
        r_lbl.font.color.rgb = BLACK
        set_font_xml(r_lbl, FONT_SEGOE)

        r_tab = p_hdr.add_run("\t")
        r_tab.font.name = FONT_SEGOE

        r_val = p_hdr.add_run(val)
        r_val.font.name = FONT_SEGOE
        r_val.font.size = Pt(9.9)
        r_val.font.bold = False
        r_val.font.color.rgb = BLACK
        set_font_xml(r_val, FONT_SEGOE)

    # Greek message body
    p_greek = doc.add_paragraph()
    p_greek.paragraph_format.space_before = Pt(20)
    p_greek.paragraph_format.space_after = Pt(12)
    p_greek.paragraph_format.line_spacing = 1.15
    r_greek = p_greek.add_run("Καλησπέρα! Το έλαβα χθες μετά τις 5 το απόγευμα")
    r_greek.font.name = FONT_APTOS
    r_greek.font.size = Pt(12)
    r_greek.font.bold = True
    r_greek.font.color.rgb = BLACK
    set_font_xml(r_greek, FONT_APTOS)

    # Mobile signature line
    p_mobile = doc.add_paragraph()
    p_mobile.paragraph_format.space_before = Pt(0)
    p_mobile.paragraph_format.space_after = Pt(6)
    p_mobile.paragraph_format.line_spacing = 1.15
    r_mob1 = p_mobile.add_run("Sent from ")
    r_mob1.font.name = FONT_APTOS
    r_mob1.font.size = Pt(12)
    r_mob1.font.color.rgb = BLACK
    set_font_xml(r_mob1, FONT_APTOS)

    r_mob2 = p_mobile.add_run("Outlook for iOS")
    r_mob2.font.name = FONT_APTOS
    r_mob2.font.size = Pt(12)
    r_mob2.font.color.rgb = BLUE_LINK
    r_mob2.font.underline = True
    set_font_xml(r_mob2, FONT_APTOS)

    # Page 1 English Screenshot
    p_crop1 = doc.add_paragraph()
    p_crop1.paragraph_format.space_before = Pt(0)
    p_crop1.paragraph_format.space_after = Pt(0)
    p_crop1.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_crop1 = p_crop1.add_run()
    r_crop1.add_picture(str(crop1_path), width=Inches(printable_width_in))

    # Page 2 English Screenshot
    doc.add_page_break()
    p_crop2 = doc.add_paragraph()
    p_crop2.paragraph_format.space_before = Pt(0)
    p_crop2.paragraph_format.space_after = Pt(0)
    p_crop2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_crop2 = p_crop2.add_run()
    r_crop2.add_picture(str(crop2_path), width=Inches(printable_width_in))

    doc.save(str(output_docx))
    print(f"✅ Document successfully created: {output_docx} ({output_docx.stat().st_size:,} bytes)")
    return output_docx


# =============================================================================
# 3. MASTER ROUTER
# =============================================================================

def create_ug100_document(
    input_pdf: str | Path,
    output_docx: Optional[str | Path] = None,
    margin_mm: float = 15.0,
    dpi: int = 300,
) -> Path:
    pdf_path = Path(input_pdf).resolve()
    if not pdf_path.exists():
        workspace = pdf_path.parent if pdf_path.parent.name not in ("input", "finished") else pdf_path.parent.parent
        for candidate in [workspace / "input" / pdf_path.name, workspace / "finished" / pdf_path.name]:
            if candidate.exists():
                pdf_path = candidate
                break
        else:
            raise FileNotFoundError(f"PDF not found: {pdf_path}")

    workspace = pdf_path.parent.parent if pdf_path.parent.name in ("input", "finished") else pdf_path.parent
    if output_docx is None:
        output_docx = workspace / "output" / f"{pdf_path.stem}.docx"
    output_docx = Path(output_docx).resolve()
    output_docx.parent.mkdir(parents=True, exist_ok=True)

    renders_dir = workspace / "renders"
    renders_dir.mkdir(parents=True, exist_ok=True)

    # Detect document type from first page text
    doc_pdf = fitz.open(str(pdf_path))
    first_page_text = doc_pdf[0].get_text()
    doc_pdf.close()

    if "ΑΝΑΦΟΡΑ" in first_page_text or "DOC-000226884" in first_page_text or "reporting form" in first_page_text:
        return build_ug100_form(pdf_path, output_docx, renders_dir, margin_mm, dpi)
    else:
        return build_ug100_email(pdf_path, output_docx, renders_dir, margin_mm, dpi)


def main():
    parser = argparse.ArgumentParser(
        description="Universal UG-100 builder for translation documents (Form & Email workflows)."
    )
    parser.add_argument("-i", "--input", required=True, help="Input PDF file path")
    parser.add_argument("-o", "--output", default=None, help="Output .docx file path")
    parser.add_argument("--margin-mm", type=float, default=15.0, help="Margin in mm (default: 15.0)")
    parser.add_argument("--dpi", type=int, default=300, help="Crop rendering DPI (default: 300)")
    args = parser.parse_args()

    create_ug100_document(
        input_pdf=args.input,
        output_docx=args.output,
        margin_mm=args.margin_mm,
        dpi=args.dpi,
    )


if __name__ == "__main__":
    main()
