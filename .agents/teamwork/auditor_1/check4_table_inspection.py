import docx
import xml.etree.ElementTree as ET

doc = docx.Document(r"output/DOC0000074469-Non-Parsable-en-US#FPREP_DXEGTK#.docx")

print("=== CHECK 4: TABLE INSPECTION & LAYOUT TABLE PROHIBITION ===")
print(f"Total tables in docx: {len(doc.tables)}")

for t_idx, table in enumerate(doc.tables):
    print(f"\n--- Table {t_idx} ---")
    print(f"Rows: {len(table.rows)}, Cols: {len(table.columns)}")
    for r_idx, row in enumerate(table.rows):
        for c_idx, cell in enumerate(row.cells):
            txt = cell.text.replace("\n", " ").strip()
            print(f"  Row {r_idx}, Col {c_idx}: '{txt[:100]}...'")

# Check all sections for w:cols
print("\n--- Sections and Multi-Column XML Configuration ---")
ns = {"w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main"}
for s_idx, section in enumerate(doc.sections):
    sectPr = section._sectPr
    cols = sectPr.findall("w:cols", ns)
    cols_info = []
    for c in cols:
        num = c.attrib.get(f"{{{ns['w']}}}num", "1")
        space = c.attrib.get(f"{{{ns['w']}}}space", "none")
        eq = c.attrib.get(f"{{{ns['w']}}}equalWidth", "none")
        col_children = c.findall("w:col", ns)
        widths = [ch.attrib.get(f"{{{ns['w']}}}w") for ch in col_children]
        cols_info.append(f"num={num}, space={space}, equal={eq}, col_widths={widths}")
    
    start = sectPr.find("w:type", ns)
    start_type = start.attrib.get(f"{{{ns['w']}}}val") if start is not None else "nextPage"
    
    lm = section.left_margin.inches
    rm = section.right_margin.inches
    tm = section.top_margin.inches
    bm = section.bottom_margin.inches
    print(f"Section {s_idx+1}: start={start_type}, cols=[{', '.join(cols_info)}], margins: L={lm:.2f}\" R={rm:.2f}\" T={tm:.2f}\" B={bm:.2f}\" (grid={8.27-lm-rm:.2f}\")")
