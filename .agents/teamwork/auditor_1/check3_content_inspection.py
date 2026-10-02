import docx
from docx.oxml import OxmlElement
from docx.shared import Pt, RGBColor
import xml.etree.ElementTree as ET

doc = docx.Document(r"output/DOC0000074469-Non-Parsable-en-US#FPREP_DXEGTK#.docx")

print("=== CONTENT GENUINENESS & SELECTIVE SCREENING INSPECTION ===")
print(f"Total paragraphs in document: {len(doc.paragraphs)}")

# 1. Inspect images and drawings across all paragraphs
images_in_doc = []
hidden_runs = []
spanish_sample_runs = []
review_stamp_found = []

ns = {
    "w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main",
    "a": "http://schemas.openxmlformats.org/drawingml/2006/main",
    "r": "http://schemas.openxmlformats.org/officeDocument/2006/relationships",
}

for p_idx, p in enumerate(doc.paragraphs):
    p_text = p.text.strip()
    
    # Check for review stamp
    if "Reviewed by TP" in p_text or "DOC0000074469" in p_text:
        review_stamp_found.append((p_idx, p_text))
        
    # Check for drawing/image in paragraph xml
    drawings = p._p.findall(".//w:drawing", ns)
    if drawings:
        for d in drawings:
            blips = d.findall(".//a:blip", ns)
            for b in blips:
                embed_id = b.attrib.get(f"{{{ns['r']}}}embed")
                # find relationship target
                rel = p.part.rels.get(embed_id)
                target = rel.target_ref if rel else "unknown"
                images_in_doc.append((p_idx, p_text[:50], target))
                
    # Check runs in paragraph
    for r_idx, r in enumerate(p.runs):
        # Check if hidden run (<w:vanish/>)
        rPr = r._r.find("w:rPr", ns)
        is_hidden = False
        if rPr is not None:
            if rPr.find("w:vanish", ns) is not None or rPr.find("w:hidden", ns) is not None:
                is_hidden = True
        if r.font.hidden or is_hidden:
            hidden_runs.append((p_idx, r_idx, len(r.text), r.text[:80]))
            
        # Sample Spanish text runs
        if any(w in r.text for w in ["ANTECEDENTES", "CASO CLÍNICO", "DISCUSIÓN", "CONCLUSIONES", "Dermatología"]):
            spanish_sample_runs.append((p_idx, r_idx, r.text, r.font.name, r.font.size, r.font.bold, r.font.italic, r.font.color.rgb if r.font.color else None))

print(f"\n--- Review Stamp ---")
for idx, text in review_stamp_found:
    print(f"Paragraph {idx}: '{text}'")

print(f"\n--- Drawings / Images in Paragraphs ({len(images_in_doc)} found) ---")
for p_idx, preview, target in images_in_doc:
    print(f"  Paragraph {p_idx} (text preview: '{preview}'): image target = {target}")

print(f"\n--- Hidden Text Runs ({len(hidden_runs)} found) ---")
for p_idx, r_idx, length, snippet in hidden_runs:
    print(f"  Paragraph {p_idx}, run {r_idx} (len={length}): '{snippet}...'")

print(f"\n--- Spanish Sample Headings / Runs ({len(spanish_sample_runs)} found) ---")
for p_idx, r_idx, text, font, size, bold, italic, color in spanish_sample_runs[:15]:
    print(f"  P{p_idx} R{r_idx}: '{text}' | font={font}, size={size}, bold={bold}, italic={italic}, color={color}")
